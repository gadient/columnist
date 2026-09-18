"""One source in, normalized text + stable locators out.

Everything here runs **before** any model call and costs nothing but CPU. That is the point of
this stage: an input we cannot read reliably is rejected with a specific, actionable code rather
than half-read and quietly analyzed.

Locators are what make evidence checkable later. A recommendation cites `paragraph:12`, and that
excerpt must be findable in the normalized note — so the locators produced here are the anchor
for "evidence before confidence".

**The document is untrusted input.** It is user-supplied, may be adversarial, and its text will
later be shown to a model. The XML parser below is hardened accordingly (no entity resolution, no
network) — that is defence against a malicious *file*, distinct from the prompt-injection defence
that protects against malicious *content* (see prompt.py).
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from lxml import etree

from .errors import ImportError

# Raw pasted or uploaded input is capped at 1 MB.
MAX_RAW_BYTES = 1024 * 1024

# Decompression-bomb defence: a DOCX under the 1 MB raw cap can still expand to
# hundreds of MB — a small zip with a highly compressible `word/document.xml`. The container reader
# below rejects on these bounds *before* decompressing anything, and the body is then read with a
# hard cap so a lying central-directory size can't slip through. Generous vs a real note (a 1 MB
# docx of text expands to a few MB), tight vs a bomb, which expands far beyond its compressed size.
_MAX_ZIP_ENTRIES = 2048          # a real docx has tens of parts, not thousands
_MAX_TOTAL_EXPANDED = 128 * 1024 * 1024
_MAX_XML_BYTES = 40 * 1024 * 1024   # word/document.xml, expanded
_MAX_COMPRESSION_RATIO = 200     # text compresses ~5–20:1; a bomb is 1000:1+

_UTF8_BOM = b"\xef\xbb\xbf"
_ZIP_MAGIC = b"PK\x03\x04"
# OLE2 compound file. Both a legacy .doc AND an *encrypted* OOXML file look like this — an
# encrypted .docx is an OLE2 wrapper around the real zip, which is why "encrypted" and "legacy"
# cannot be told apart by magic bytes alone. We route on the declared extension instead.
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Revision markup. Their presence means the visible text depends on whether you accept or reject —
# so reading it would be a guess. We reject rather than guess.
_TRACKED_CHANGE_TAGS = ("ins", "del", "moveFrom", "moveTo")

# Body content we do not read. Text is still extracted; the user is told what was skipped.
_NON_TEXT_TAGS = {
    "drawing": "images or drawings",
    "pict": "pictures",
    "object": "embedded objects",
    "txbxContent": "text boxes",
}

# Hardened parser: the DOCX is untrusted. No entity resolution (billion-laughs / XXE), no network
# fetches for external entities, no unbounded tree.
_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)


@dataclass(frozen=True)
class Segment:
    """One locatable unit of the note. `locator` is stable and is what evidence cites."""

    text: str
    locator: str  # "line:5" | "paragraph:12" | "table:1:row:3:cell:2"


@dataclass(frozen=True)
class Extraction:
    text: str  # the normalized note handed to the model — never truncated
    segments: tuple[Segment, ...]
    warnings: tuple[str, ...]  # things present but not analyzed


def _normalize_newlines(text: str) -> str:
    """Line endings to \\n for analysis; source line numbers stay stable because we only ever
    collapse a 2-char terminator to 1, never add or remove lines."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _from_lines(text: str) -> Extraction:
    normalized = _normalize_newlines(text)
    segments = tuple(
        Segment(line.strip(), f"line:{i}")
        for i, line in enumerate(normalized.split("\n"), start=1)
        if line.strip()
    )
    if not segments:
        # Whitespace-only paste, or a file of blank lines.
        raise ImportError("EMPTY_INPUT")
    return Extraction(text=normalized, segments=segments, warnings=())


def extract_pasted(text: str) -> Extraction:
    """Pasted text: already a Unicode string from the browser, so there is no encoding to guess."""
    return _from_lines(text)


def extract_txt(raw: bytes) -> Extraction:
    """UTF-8 and UTF-8-with-BOM only. We do **not** guess at any other encoding.

    A Latin-1 file containing only ASCII decodes cleanly here — correctly, because such a file
    *is* valid UTF-8. Only genuinely non-UTF-8 bytes are rejected.
    """
    body = raw[len(_UTF8_BOM) :] if raw.startswith(_UTF8_BOM) else raw
    try:
        text = body.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ImportError(
            "UNSUPPORTED_ENCODING",
            detail=f"Not valid UTF-8 (byte {exc.start} of the file).",
        ) from exc
    return _from_lines(text)


def _reject_bad_container(raw: bytes, *, declared_ext: str) -> zipfile.ZipFile:
    """Everything that makes a file not-a-readable-DOCX, in order of cheapness."""
    if raw[:8] == _OLE2_MAGIC:
        # A real .doc, or an encrypted .docx. Distinguish on what the user claimed it was.
        if declared_ext == ".doc":
            raise ImportError(
                "UNSUPPORTED_FILE_TYPE",
                detail="This is a legacy Word (.doc) file. Only .docx is supported.",
            )
        raise ImportError(
            "DOCX_INVALID",
            detail="This file is encrypted, or is a legacy .doc renamed to .docx.",
        )
    if raw[:4] != _ZIP_MAGIC:
        # A .txt/.pdf/anything renamed .docx, or a truncated file.
        raise ImportError("DOCX_INVALID", detail="This is not a DOCX file.")
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as exc:
        raise ImportError("DOCX_INVALID", detail="The file is corrupt and cannot be opened.") from exc

    # Decompression-bomb bounds, from the central directory — cheap, no decompression yet.
    infos = zf.infolist()
    if len(infos) > _MAX_ZIP_ENTRIES:
        raise ImportError("DOCX_INVALID", detail="This file has too many internal parts.")
    total_expanded = sum(i.file_size for i in infos)
    if total_expanded > _MAX_TOTAL_EXPANDED:
        raise ImportError("DOCX_INVALID", detail="This file expands to an implausible size.")
    compressed = sum(i.compress_size for i in infos) or 1
    if total_expanded / compressed > _MAX_COMPRESSION_RATIO:
        raise ImportError("DOCX_INVALID", detail="This file's compression ratio looks malicious.")

    names = set(zf.namelist())
    if any(n.startswith("word/vbaProject") for n in names):
        # .docm wearing a .docx name, or a genuine .docm. We do not read macro-enabled documents.
        raise ImportError(
            "UNSUPPORTED_FILE_TYPE",
            detail="This is a macro-enabled document (.docm). Save it as a plain .docx.",
        )
    if "word/document.xml" not in names:
        # A valid zip that is not a Word document (e.g. .xlsx, or a plain archive renamed .docx).
        raise ImportError("DOCX_INVALID", detail="This is not a DOCX file.")
    return zf


def _read_capped(zf: zipfile.ZipFile, name: str, cap: int) -> bytes:
    """Decompress at most `cap` bytes of a zip entry, so a lying central-directory `file_size`
    cannot slip a bomb past the metadata checks: `ZipExtFile.read(n)` decompresses on demand, so
    reading `cap + 1` never expands the whole entry into memory."""
    try:
        with zf.open(name) as f:
            data = f.read(cap + 1)
    except KeyError as exc:  # pragma: no cover - guarded by _reject_bad_container
        raise ImportError("DOCX_INVALID", detail="This is not a DOCX file.") from exc
    if len(data) > cap:
        raise ImportError("DOCX_INVALID", detail="The document body is too large to read.")
    return data


def _reject_tracked_changes(root: etree._Element) -> None:
    for tag in _TRACKED_CHANGE_TAGS:
        if root.find(f".//{_W}{tag}") is not None:
            raise ImportError("DOCX_TRACKED_CHANGES")


def _warn_non_text(root: etree._Element) -> tuple[str, ...]:
    found = [
        label for tag, label in _NON_TEXT_TAGS.items() if root.find(f".//{_W}{tag}") is not None
    ]
    if not found:
        return ()
    return (
        "This document contains "
        + ", ".join(found)
        + ". Only paragraph and table text was analyzed.",
    )


def extract_docx(raw: bytes, *, declared_ext: str = ".docx") -> Extraction:
    """Body paragraphs and table cells, in document order, with stable locators.

    Deliberately not read: headers, footers, comments, footnotes, endnotes. Those live in separate
    parts of the container (``word/header1.xml`` and friends) and we simply never open them — the
    exclusion is structural, not a filter that could be forgotten.
    """
    zf = _reject_bad_container(raw, declared_ext=declared_ext)
    # Capped, on-demand read — never trusts the declared size (see _read_capped). NOTE: python-docx
    # below re-opens the package itself; the central-directory bounds in _reject_bad_container are
    # what guard that second parse. A fully bounded reader for every part is a follow-up.
    document_xml = _read_capped(zf, "word/document.xml", _MAX_XML_BYTES)

    try:
        root = etree.fromstring(document_xml, parser=_XML_PARSER)
    except etree.XMLSyntaxError as exc:
        raise ImportError("DOCX_INVALID", detail="The document's contents are corrupt.") from exc

    _reject_tracked_changes(root)
    warnings = _warn_non_text(root)

    # python-docx does the text assembly (runs, tabs, breaks); we drive the document-order walk
    # ourselves because `doc.paragraphs` and `doc.tables` are separate lists and lose interleaving.
    from docx import Document  # lazy: only imported on the DOCX path
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(raw))
    except Exception as exc:  # python-docx raises assorted types on malformed parts
        raise ImportError("DOCX_INVALID", detail="The document could not be read.") from exc

    segments: list[Segment] = []
    para_index = 0
    table_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            # Count every paragraph, emit only non-empty ones: locators stay true to the document
            # even where blank paragraphs are used for spacing.
            para_index += 1
            text = Paragraph(child, doc).text.strip()
            if text:
                segments.append(Segment(text, f"paragraph:{para_index}"))
        elif child.tag == qn("w:tbl"):
            table_index += 1
            for row_i, row in enumerate(Table(child, doc).rows, start=1):
                for cell_i, cell in enumerate(row.cells, start=1):
                    text = cell.text.strip()
                    if text:
                        segments.append(
                            Segment(text, f"table:{table_index}:row:{row_i}:cell:{cell_i}")
                        )

    if not segments:
        # A document of only images/scans, or genuinely empty. Distinct from a corrupt file —
        # we read it fine, there was just nothing to read.
        raise ImportError("NO_EXTRACTABLE_TEXT")

    return Extraction(
        text="\n".join(s.text for s in segments),
        segments=tuple(segments),
        warnings=warnings,
    )
